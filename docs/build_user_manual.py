from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path(__file__).parent/'Customer_Signals_Customer_User_Manual.docx'
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.7); s.bottom_margin=Inches(.7); s.left_margin=Inches(.8); s.right_margin=Inches(.8)
d.styles['Normal'].font.name='Aptos'; d.styles['Normal'].font.size=Pt(10.5); d.styles['Normal'].font.color.rgb=RGBColor(35,45,48)
d.styles['Normal'].paragraph_format.space_after=Pt(5); d.styles['Normal'].paragraph_format.line_spacing=1.1
for name,size in [('Title',29),('Subtitle',13),('Heading 1',18),('Heading 2',13)]:
 st=d.styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0); st.font.bold=name!='Subtitle'; st.paragraph_format.keep_with_next=True
d.styles['Heading 1'].paragraph_format.space_before=Pt(15); d.styles['Heading 1'].paragraph_format.space_after=Pt(6)
header=s.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('CUSTOMER SIGNALS  /  CUSTOMER USER MANUAL'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(85,100,102)
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Customer Signals  •  Page '); r.font.size=Pt(8); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
def p(x='',style=None): return d.add_paragraph(x,style)
def h(x,l=1): d.add_heading(x,l)
def b(x): d.add_paragraph(x,'List Bullet')
def n(x): d.add_paragraph(x,'List Number')
def tbl(head,rows,widths):
 t=d.add_table(rows=1,cols=len(head)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for i,x in enumerate(head):
  c=t.rows[0].cells[i]; c.text=x; sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'24495A'); c._tc.get_or_add_tcPr().append(sh)
  for r in c.paragraphs[0].runs: r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(255,255,255)
 repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  for i,x in enumerate(row):
   c=cells[i]; c.text=str(x)
   if ri%2: sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F2F6F7'); c._tc.get_or_add_tcPr().append(sh)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(1)
    for r in q.runs: r.font.size=Pt(8.5); r.font.name='Aptos'
 borders=OxmlElement('w:tblBorders')
 for edge in ('top','left','bottom','right','insideH','insideV'):
  e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'D9E0E2'); borders.append(e)
 t._tbl.tblPr.append(borders)
 for row in t.rows:
  for c,w in zip(row.cells,widths): c.width=Inches(w)
 d.add_paragraph().paragraph_format.space_after=Pt(1)

x=p(style='Title'); x.paragraph_format.space_before=Pt(105); x.add_run('Customer Signals')
p('Customer User Manual','Subtitle')
x=p('A practical guide to customer profiles, segments, recommendations, support, evaluation, and exports.'); x.paragraph_format.space_before=Pt(20); x.runs[0].font.size=Pt(15)
x=p('Version 1.0  |  28 October 2026'); x.paragraph_format.space_before=Pt(24); x.runs[0].font.color.rgb=RGBColor(75,95,101)
x=p('For authorized workspace users'); x.paragraph_format.space_before=Pt(180)
d.add_page_break()

h('1  About this manual')
p('This guide explains how authorized staff use the Customer Signals workspace. The screens and steps describe the current application; your account permissions determine which pages and actions are available.')
p('Customer Signals organizes customer, purchase, engagement, consent, support, segment, and action data. It generates explainable segments and recommendations and records delivery and conversion outcomes. Users remain responsible for reviewing configuration, consent, and delivery readiness before contacting customers.')
h('Operating boundaries',2)
b('Customer profiles are read views. Source records and immutable history are not edited from the profile.')
b('Email dispatch is a controlled server-side operation, not a workspace button. Default operation is a preview; actual sending needs authorization and SMTP configuration.')
b('The current delivery adapter supports email only. SMS and push recommendations are not sent by a configured provider.')
b('Demo data generation is for isolated development/acceptance environments, never a production database.')

h('2  Roles and access')
p('Administrators provision individual accounts and assign least-privilege groups. An account may have multiple groups only when approved.')
tbl(['Role','Typical work','Boundary'],[
('System Administrators','Manage users, groups, rules, actions, products, and app configuration.','Django admin access; restrict to trusted administrators.'),
('Marketing Operations','Review campaigns; rebuild segments; run recommendations; update outcomes; export permitted data.','Workflow access does not authorize real email delivery.'),
('Support Agents','Review customer context; open, resolve, reopen cases.','Support permissions; no workflow execution by default.'),
('Analysts','Read permitted records and evaluation; export authorized data.','Read-oriented; no mutations by default.')],[1.35,2.7,2.45])
p('Sign-in is required. The admin console also requires staff status and model permissions. If an item or action is missing, ask an administrator to verify your role; do not use another person’s account.')

h('3  Sign in and navigate')
h('Sign in',2)
n('Open the workspace address supplied by your organization. 127.0.0.1:8000 is a local development address; production should use its approved HTTPS site.')
n('Enter your provisioned username and password. Authentication uses Django’s username field; some deployments use email as the username, so confirm with your administrator.')
n('After sign-in, the workspace opens Overview. If login fails, check case and spelling; ask an administrator to confirm the account is active and correctly assigned.')
h('Navigation',2)
b('Overview summarizes customers, segment mix, recent runs, and—when permitted—action performance.')
b('Customers opens the searchable directory. Select a name or arrow to open the profile.')
b('Evaluation shows action metrics, the journey chart, action log, permitted outcome updates, and Export action log.')
b('Action log links to the log section on Evaluation. Admin console opens Django administration for staff users.')
b('Select Customer Signals in the admin header to return to the workspace dashboard. The theme switch changes appearance and saves locally in that browser.')
b('Use Sign out in the workspace footer when finished.')

h('4  Overview and run monitoring')
n('Review the Customers card and Current segment mix. Rows count current memberships under active rules; bars show share relative to all customers.')
n('If permitted, review Action records, Conversion rate, Attributed revenue, and Suppressed. Conversion rate uses sent actions; suppressed records are excluded from its denominator.')
n('Check Workflow activity. A queued job may display “Waiting for worker” until a worker starts it.')
n('Review Recent recommendations for customer, action, time, status, and suppression reason. Select a customer ID to open the profile.')
tbl(['State','Meaning','What to do'],[('Queued','Accepted, waiting to run.','Wait, then refresh Overview.'),('Running','Worker started.','Allow completion; avoid duplicate submissions.'),('Succeeded','Completed with a summary.','Review segment mix or action log.'),('Failed','Did not complete.','Share run ID and visible error with the operator.')],[1.05,2.55,2.9])

h('5  Customers and profiles')
h('Find a customer',2)
n('Open Customers and search by customer ID, email, or surname. Partial matches are supported.')
n('Select Search; Clear removes the query. Results are paginated at 50 per page with Previous and Next controls.')
n('Select the customer name or arrow to open the profile.')
h('Read a profile',2)
tbl(['Panel','Information'],[
('Customer details','ID, email, region, join date, lifecycle.'),('Latest orders','Up to 20 recent orders with dates, state, and total. Returned/cancelled orders are visible but excluded from purchase RFM.'),('Segment history','Up to 10 recent decisions, score, explanation, time, and current marker; earlier records remain historical.'),('Action history','Up to 10 recent recommendations, channel, state, suppression reason (if permitted).'),('Channel consent','Email/SMS/push decisions and effective dates. Email consent is checked again before delivery.'),('Support cases','Up to 20 cases with category, state, cost, description, and event history.')],[1.45,5.05])
p('A “No … recorded” message means the profile has no matching data. Report suspected source-data errors with the customer ID and field to the data owner.')

h('6  Support case workflow')
p('Support Agents open cases from a customer profile. Accepted case state changes are kept in an event history with actor and time.')
h('Open a case',2)
n('Find the customer. In Support cases, choose Open case. If absent, request permission to add cases.')
n('Select Billing, Delivery, Product, or Other. Enter a factual description and service cost if known; blank cost becomes zero.')
n('Submit, note the case ID, and confirm it appears in the profile.')
h('Resolve or reopen',2)
n('Choose Update case beside the case. Open cases can be resolved; resolved cases can be reopened.')
n('Select the available state, write concise notes, and submit. Reopening requires an explanation.')
n('Confirm the state and history event appear. Cases cannot be deleted through this workflow.')
p('Service cost of $100 or more may make a customer eligible for Service Attention after a rebuild. Avoid unnecessary sensitive data in notes.')

h('7  Segmentation')
p('Segmentation assigns current memberships from purchase, service, and engagement signals. Each decision stores a rule snapshot, score, explanation, and time. Previous membership records remain as history.')
h('Prepare and rebuild',2)
n('Confirm customer, order, event, support, and segment-rule records are current; relevant rules must be active and effective.')
n('On Overview choose Rebuild segments. This control requires segmentation permission.')
n('Record the run ID and monitor Workflow activity to Succeeded or Failed. Do not submit the same run repeatedly while it is processing.')
n('Review Current segment mix and spot-check customer explanations.')
tbl(['Priority','Segment','Service condition'],[('1','New','Customer classified as new.'),('2','Dormant','No completed purchase in 180+ days.'),('3','High Value','$1,000+ net completed spend across 3+ completed orders.'),('4','Service Attention','Support-case service cost is $100 or more.'),('5','Engaged','2+ engagement events in the last 30 days.'),('6','Active','Fallback when no higher-priority rule matches.')],[.7,1.55,4.25])
p('Returns and cancellations do not contribute to RFM calculations. The profile’s recorded reason is the best explanation of a specific assignment.')

h('8  Recommendations and delivery')
p('Recommendation runs match customers to active actions. Eligibility checks segment membership, effective dates, channel consent, and cooldown. Suppressed recommendations remain visible with a reason.')
n('Review segmentation and active action configuration first.')
n('On Overview choose Run recommendations; record the run ID and monitor completion.')
n('Review new records in Evaluation or profiles. Check suppression reasons; do not bypass them.')
p('Cooldown is configured in days for each action. A recommendation is not proof that a message was sent.')
h('Email delivery control',2)
p('Only an authorized deployment operator runs the email dispatcher. Without --send it is preview-only and reports eligible, suppressed, failed, and non-email records. Actual sending requires --send, ACTION_DELIVERY_ENABLED=true, approved SMTP configuration/content/sender, and a fresh consent check. Keep sending disabled until the organization approves it.')
p('SMS and push remain queued because provider adapters are not configured. Do not report them as delivered.')

h('9  Evaluation and outcomes')
h('Review results',2)
n('Open Evaluation and review Recommendations, Sent, Conversions, Attributed revenue, and Suppressed.')
n('Use Action journey to compare counts. The Action log lists 50 records per page; use pagination for older records.')
n('Select the customer ID for context. Select Update only if your permission and the record state allow it.')
p('Conversion rate is conversions divided by sent. Suppressed records do not count as sent. Revenue comes from recorded conversions. Labeled synthetic acceptance anomalies are excluded from metrics; do not interpret them as real activity.')
h('Record an outcome',2)
n('For Recommended, choose Update and select Sent or Suppressed. Consent is validated before Sent is accepted.')
n('For Sent, choose Converted or Not converted. Converted requires nonnegative attributable revenue.')
n('Add concise evidence notes and submit. Confirm the log and metrics update.')
p('Valid transitions: Recommended → Sent or Suppressed; Sent → Converted or Not converted. Completed states are terminal. Every accepted change records actor and time. Mark Sent only when provider acceptance or approved evidence confirms delivery.')

h('10  Export CSV data')
p('Exports are streamed, ordered CSV downloads with permission checked for each type. A link may be hidden if access is not granted.')
tbl(['File','Where','Contents'],[('customers.csv','Customers → Export customers','Customer identifiers and profile/contact fields.'),('memberships.csv','Overview → Export','Membership IDs, customer/segment IDs, score, current state, reason, rule snapshot.'),('actions.csv','Evaluation → Export action log','Outcome IDs, customer/action references, dates, status, suppression, revenue, action snapshot.')],[1.2,2.25,3.05])
n('Choose the relevant Export control and allow the browser download.')
n('Open the CSV with an approved tool. Preserve IDs as text and check timestamp interpretation.')
n('Store/share only under customer-data policy; remove temporary copies when no longer needed.')
p('Files may contain personal or commercially sensitive information. CSV formula-injection protections do not remove the need to treat values as untrusted. If export fails, report its type, page, and time.')

h('11  Admin console')
p('Django admin is for authorized configuration and inspection. Open Admin console from the workspace as a staff user with permissions; use your provisioned username/password; select Customer Signals in the admin header to return to Overview.')
p('Depending on access, administrators can manage users/groups, products, segment rules, and actions and inspect customers, orders, events, consent, memberships, cases, outcomes, and workflow history. Some records are intentionally read-only or cannot be deleted. Use approved workflows rather than bypassing lifecycle controls.')
b('The setup_roles command creates/updates the four standard groups; an administrator assigns named users to the correct group.')
b('Use unique credentials and individual accounts. Never store passwords in this manual or share team accounts.')
b('Review rule priority/effective dates and action channel/cooldown/content with the business owner before runs.')
b('Permission changes can expose data or enable changes; follow access approval procedures.')

h('12  Troubleshooting')
tbl(['Issue','Likely reason','Response'],[
('Invalid login','Wrong username/password, inactive account; admin may require staff access.','Confirm username field/account with administrator; reset only through approved process.'),
('Missing menu or access denied','Required model/workflow permission absent.','Ask administrator to check intended group and access.'),
('Run stays Queued','Worker/broker delayed or unavailable.','Give run ID to operator to inspect queue health.'),
('Run Failed','Data/configuration, database, or worker error.','Share run ID/error; operator checks logs before retry.'),
('No customer results','Query mismatch or empty data.','Clear query, check spelling, confirm data load.'),
('Suppressed action','Consent missing/withdrawn or cooldown.','Read reason; never override consent; confirm preference with owner.'),
('Export problem','Permission, server, or browser download behavior.','Check downloads and report export type/time/page.'),
('Theme differs on device','Choice is stored in local browser storage.','Set theme on that browser; settings may not sync.'),
('Email delivery issue','Consent, non-email channel, SMTP, provider, or disabled flag.','Review dry-run/provider logs with operator; keep sending disabled until cleared.')],[1.35,2.05,3.1])

h('13  Operating checklists')
h('Marketing Operations',2)
for x in ['Confirm source data and active rules/actions.','Rebuild segments and monitor completion.','Review mix and sample explanations.','Run recommendations and inspect suppressions.','Use only the approved email preview/send procedure.','Record evidenced outcomes; enter revenue only when attributable.','Export only for approved recipients.']: b(x)
h('Support Agents',2)
for x in ['Verify the profile before opening a case.','Choose the right category and write factually.','Resolve when service work is complete.','Explain every reopen and review event history.','Keep unnecessary sensitive information out of notes.']: b(x)
h('System Administrators',2)
for x in ['Provision named accounts and least-privilege roles.','Protect customer exports and admin access.','Coordinate campaign rules and dates with business owners.','Escalate database, queue, SMTP, and deployment issues.','Never run demo generation/clearing on production.']: b(x)

h('14  Glossary and help')
tbl(['Term','Definition'],[('Action catalog','Configured action with channel, segment, effective dates, priority, and cooldown.'),('Action outcome','Customer recommendation and its lifecycle state.'),('Attributed revenue','Revenue recorded for a converted action.'),('Cooldown','Minimum wait before an action can be recommended again.'),('Current membership','Segment assigned by the latest successful rebuild.'),('Effective-dated consent','Channel permission with start and optional end; applicable decision is checked.'),('Suppressed','Recommendation held by consent or cooldown.'),('Workflow run','Tracked segmentation or recommendation batch.'),('CSV','Comma-separated file for download and analysis.')],[1.7,4.8])
p('For field definitions, use the application owner’s data dictionary. For access, data corrections, queue health, delivery setup, or deployment issues, contact your organization’s Customer Signals administrator or operations owner.')
d.core_properties.title='Customer Signals Customer User Manual'; d.core_properties.subject='User workflows and actions in Customer Signals'; d.core_properties.author='Customer Signals'
d.save(out); print(out)
